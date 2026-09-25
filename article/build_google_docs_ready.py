#!/usr/bin/env python3
"""Build Google-Docs-importable DOCX files with repository figures embedded."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "article" / "drafts"
OUTPUT_DIR = ROOT / "article" / "google_docs_ready"
SOURCES = (
    "PART_1_DATA_COLLECTION_AND_EDA.md",
    "PART_2_LLM_FEATURE_ENGINEERING.md",
    "PART_3_MODEL_EVALUATION.md",
    "PART_4_GRAPHRAG_EXPLANATIONS.md",
)

INK = "263238"
MUTED = "667085"
NAVY = "3977A8"
LIGHT = "EEF2F6"
GRID = "D9DEE5"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID, size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    run.font.name = "Arial"; run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(MUTED)


def style_document(document: Document, short_title: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.13

    for name, size, color in [
        ("Title", 22, INK), ("Heading 1", 17, INK), ("Heading 2", 13, NAVY),
    ]:
        style = document.styles[name]
        style.font.name = "Arial"; style.font.size = Pt(size); style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.text = short_title
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"; run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_number(section.footer.paragraphs[0])


INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")


def add_inline(paragraph, text: str) -> None:
    position = 0
    for match in INLINE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2]); run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1]); run.font.name = "Consolas"; run.font.size = Pt(9)
        else:
            run = paragraph.add_run(token[1:-1]); run.italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_table(document: Document, lines: list[str]) -> None:
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    columns = max(map(len, rows))
    table = document.add_table(rows=len(rows), cols=columns)
    table.autofit = True
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        row_properties = table.rows[row_index]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_properties.append(cant_split)
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = values[column_index] if column_index < len(values) else ""
            cell.text = ""
            add_inline(cell.paragraphs[0], value)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            # Keep compact publication tables together when the complete table
            # fits on one page; this prevents a single orphaned row after import.
            cell.paragraphs[0].paragraph_format.keep_with_next = row_index < len(rows) - 1
            for run in cell.paragraphs[0].runs:
                run.font.name = "Arial"; run.font.size = Pt(8.2 if columns >= 5 else 9)
                if row_index == 0:
                    run.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
            shade(cell, NAVY if row_index == 0 else (LIGHT if row_index % 2 == 0 else "FFFFFF"))
            set_cell_border(cell)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_image(document: Document, source: Path, alt: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run()
    run.add_picture(str(source), width=Inches(6.75))
    # Store meaningful fallback text in the drawing description for accessibility/import.
    drawing = run._r.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("descr", alt)


def parse_document(source: Path) -> Document:
    document = Document()
    raw_lines = source.read_text().splitlines()
    title_text = next((line[2:] for line in raw_lines if line.startswith("# ")), source.stem)
    style_document(document, title_text)
    document.core_properties.title = title_text
    document.core_properties.subject = "Technical restaurant-recommendation study"

    index = 0
    pending: list[str] = []

    def flush_pending() -> None:
        if not pending:
            return
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(value.strip() for value in pending))
        pending.clear()

    while index < len(raw_lines):
        line = raw_lines[index]
        stripped = line.strip()
        if not stripped:
            flush_pending(); index += 1; continue
        if stripped.startswith("|" ):
            flush_pending(); table_lines = []
            while index < len(raw_lines) and raw_lines[index].strip().startswith("|"):
                table_lines.append(raw_lines[index].strip()); index += 1
            add_table(document, table_lines); continue
        if stripped == "<!-- PAGEBREAK -->":
            flush_pending()
            document.add_page_break()
            index += 1
            continue
        image_match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            flush_pending()
            image_path = (source.parent / image_match.group(2)).resolve()
            if not image_path.exists():
                raise FileNotFoundError(f"Missing image for {source.name}: {image_path}")
            add_image(document, image_path, image_match.group(1)); index += 1; continue
        if stripped.startswith("# "):
            flush_pending(); paragraph = document.add_paragraph(style="Title")
            add_inline(paragraph, stripped[2:]); index += 1; continue
        if stripped.startswith("## "):
            flush_pending(); paragraph = document.add_paragraph(style="Heading 1")
            add_inline(paragraph, stripped[3:]); index += 1; continue
        if stripped.startswith("### "):
            flush_pending(); paragraph = document.add_paragraph(style="Heading 2")
            add_inline(paragraph, stripped[4:]); index += 1; continue
        if stripped.startswith("- "):
            flush_pending(); paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:]); index += 1; continue
        if stripped.startswith(">"):
            flush_pending(); quote_lines = []
            while index < len(raw_lines) and raw_lines[index].strip().startswith(">"):
                quote_lines.append(raw_lines[index].strip().lstrip(">").strip()); index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.right_indent = Inches(0.18)
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(8)
            add_inline(paragraph, " ".join(quote_lines))
            for run in paragraph.runs:
                run.italic = True; run.font.color.rgb = RGBColor.from_string(MUTED)
            continue
        if stripped.startswith("*Figure") and stripped.endswith("*"):
            flush_pending(); paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            run = paragraph.add_run(stripped[1:-1]); run.italic = True
            run.font.size = Pt(8.7); run.font.color.rgb = RGBColor.from_string(MUTED)
            index += 1; continue
        pending.append(stripped); index += 1
    flush_pending()
    return document


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for filename in SOURCES:
        source = SOURCE_DIR / filename
        output = OUTPUT_DIR / source.with_suffix(".docx").name
        document = parse_document(source)
        document.save(output)
        print(f"Wrote {output}")


if __name__ == "__main__":
    main()
